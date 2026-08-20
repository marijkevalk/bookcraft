"""book-formatter CLI."""

from __future__ import annotations

import logging
import re
from pathlib import Path

import click
from docx import Document

from bookcraft import __version__
from bookcraft.ai_chapters import (
    BACKENDS,
    DEFAULT_BACKEND,
    DEFAULT_CLAUDE_BIN,
    DEFAULT_GEMINI_MODEL,
    detect_chapters_ai,
    detect_sneak_preview,
)
from bookcraft.amazon_reviews import fetch_reviews, format_reviews_report
from bookcraft.analyze import analyze_manuscript, claude_runner, split_chapters
from bookcraft.formatter import render
from bookcraft.history import (
    build_scorecard,
    load_history,
    record_from_analysis,
    save_record,
)
from bookcraft.learning import load_learning, save_learning, update_learning
from bookcraft.metadata import parse_metadata_file
from bookcraft.metrics import ChapterMetrics, compute_metrics
from bookcraft.report import write_pdf
from bookcraft.template_builder import build_docxtpl_template
from bookcraft.verify import default_report_path, format_report, verify


@click.group()
@click.version_option(version=__version__)
def cli() -> None:
    """Format manuscript .docx files into a styled template."""


@cli.command("build-template")
@click.option(
    "--source",
    "source_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Original styled template .docx",
)
@click.option(
    "--output",
    "output_path",
    required=True,
    type=click.Path(dir_okay=False, path_type=Path),
    help="Where to write the docxtpl-ready template .docx",
)
@click.option(
    "--paperback",
    is_flag=True,
    default=False,
    help="Build a paperback template (TOC with page numbers, footer page numbers)",
)
def build_template_cmd(source_path: Path, output_path: Path, paperback: bool) -> None:
    """Convert a styled template to a docxtpl-ready template (one-time setup)."""
    build_docxtpl_template(source_path, output_path, ebook=not paperback)
    click.echo(f"Wrote {output_path}")


@cli.command()
@click.option(
    "--template",
    "template_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Original styled template .docx (the source, not the pre-built one)",
)
@click.option(
    "--source",
    "source_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Path to the manuscript .docx",
)
@click.option(
    "--metadata",
    "metadata_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Path to the book metadata .txt (title, subtitle, series, pen name)",
)
@click.option(
    "--output",
    "output_path",
    required=True,
    type=click.Path(dir_okay=False, path_type=Path),
    help="Where to write the ebook .docx (paperback gets _paperback suffix)",
)
@click.option(
    "--backend",
    type=click.Choice(BACKENDS),
    default=DEFAULT_BACKEND,
    show_default=True,
    help="Chapter-detection backend: claude (local CLI), gemini (own API "
    "key, GEMINI_API_KEY), or heuristic (no AI/account).",
)
@click.option(
    "--model",
    default=None,
    help=f"Model name for the gemini backend (default: {DEFAULT_GEMINI_MODEL}).",
)
@click.option(
    "--claude-bin",
    default=DEFAULT_CLAUDE_BIN,
    show_default=True,
    help="Claude Code CLI binary to invoke for the claude backend",
)
def format(
    template_path: Path,
    source_path: Path,
    metadata_path: Path,
    output_path: Path,
    backend: str,
    model: str | None,
    claude_bin: str,
) -> None:
    """Merge manuscript + metadata into the template using AI chapter detection.

    Produces two outputs:
    - <output>.docx           — ebook (no page numbers)
    - <output>_paperback.docx — paperback (TOC + footer page numbers)

    Chapter structure is detected by the chosen ``--backend``: the local
    Claude Code CLI (default), Google's Gemini API with your own key, or a
    zero-AI heuristic.
    """
    import tempfile

    log_path = output_path.with_suffix(".log")
    logging.basicConfig(
        level=logging.DEBUG,
        format="%(asctime)s %(levelname)-8s %(name)s: %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8")],
    )
    click.echo(f"Logging to {log_path}")

    click.echo(f"Reading {source_path.name}...")
    metadata = parse_metadata_file(metadata_path)
    source_doc = Document(str(source_path))

    click.echo(f"Detecting chapters ({backend} backend)...")
    chapters = detect_chapters_ai(
        source_doc, backend=backend, model=model, claude_bin=claude_bin
    )

    if not chapters:
        raise click.ClickException(f"No chapters detected in {source_path}")

    click.echo(
        f"Found {len(chapters)} chapters: {', '.join(c.title for c in chapters)}"
    )

    sneak_preview = detect_sneak_preview(source_doc)
    if sneak_preview:
        click.echo(f"Found sneak preview ({len(sneak_preview)} paragraphs).")

    paperback_path = output_path.with_name(
        output_path.stem + "_paperback" + output_path.suffix
    )

    chapter_summary = (
        f"{len(chapters)} chapters: {', '.join(c.title for c in chapters)}"
    )

    click.echo("Rendering ebook and paperback...")
    with tempfile.TemporaryDirectory() as tmp:
        ebook_tpl = Path(tmp) / "template-ebook.docx"
        pb_tpl = Path(tmp) / "template-paperback.docx"
        build_docxtpl_template(template_path, ebook_tpl, ebook=True)
        build_docxtpl_template(template_path, pb_tpl, ebook=False)
        render(
            ebook_tpl, metadata, chapters, output_path, sneak_preview_body=sneak_preview
        )
        render(
            pb_tpl, metadata, chapters, paperback_path, sneak_preview_body=sneak_preview
        )

    click.echo(f"Wrote {output_path} ({chapter_summary})")
    click.echo(f"Wrote {paperback_path} ({chapter_summary})")


@cli.command("verify")
@click.option(
    "--source",
    "source_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Manuscript .docx (the source you ran `format` against)",
)
@click.option(
    "--output",
    "output_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Rendered output .docx produced by `format`",
)
@click.option(
    "--report",
    "report_path",
    required=False,
    default=None,
    type=click.Path(dir_okay=False, path_type=Path),
    help=(
        "Where to write the verification report .txt "
        "(defaults to <output>_verification.txt next to the output file)"
    ),
)
def verify_cmd(source_path: Path, output_path: Path, report_path: Path | None) -> None:
    """Verify the rendered output matches the source manuscript verbatim.

    Compares text and italic-run positions per body paragraph. Writes a
    full per-paragraph report (with character counts) to a .txt file and
    prints a one-line summary. Exits non-zero if any mismatches are found.
    """
    result = verify(source_path, output_path)
    report_target = report_path or default_report_path(output_path)
    report_target.write_text(
        format_report(result, source_path=source_path, output_path=output_path),
        encoding="utf-8",
    )
    click.echo(
        f"Body paragraphs: {result.source_para_count} source, "
        f"{result.output_para_count} output"
    )
    click.echo(
        f"Total characters: {result.total_source_chars:,} source, "
        f"{result.total_output_chars:,} output"
    )
    click.echo(
        f"Text mismatches:   {result.text_mismatches} / {len(result.all_paragraphs)}"
    )
    click.echo(
        f"Italic mismatches: {result.italic_mismatches} / {len(result.all_paragraphs)}"
    )
    click.echo(f"Report written to: {report_target}")
    if not result.content_clean:
        raise click.ClickException(
            "Verification found content mismatches; see report for details."
        )
    if not result.structure_clean:
        click.echo(
            f"Warning: chapter count differs ({result.source_chapter_count} source "
            f"vs {result.output_chapter_count} output) — Claude may have missed a "
            "heading. Check the output manually."
        )


@cli.command("fetch-reviews")
@click.argument("asin")
@click.option(
    "--output",
    "output_path",
    required=False,
    default=None,
    type=click.Path(dir_okay=False, path_type=Path),
    help="Where to write the reviews .txt (default: <asin>_reviews.txt in current dir)",
)
@click.option(
    "--pages",
    default=5,
    show_default=True,
    help="Number of review pages to fetch (10 reviews per page)",
)
def fetch_reviews_cmd(asin: str, output_path: Path | None, pages: int) -> None:
    """Fetch Amazon customer reviews for a book by ASIN.

    ASIN is the Amazon product ID, e.g. B0GX2ZW64X.
    Writes a plain-text report with all reviews grouped by star rating.
    """
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    click.echo(f"Fetching reviews for ASIN {asin} (up to {pages} pages)...")
    reviews = fetch_reviews(asin, max_pages=pages)

    report = format_reviews_report(reviews, asin)

    target = output_path or Path(f"{asin}_reviews.txt")
    target.write_text(report, encoding="utf-8")

    click.echo(f"Fetched {len(reviews)} reviews → {target}")


def _extract_asin(book: str) -> str:
    """Pull a 10-char ASIN out of an Amazon URL, or accept a bare ASIN."""
    book = book.strip()
    match = re.search(r"(?:/dp/|/gp/product/|/product/)([A-Z0-9]{10})", book)
    if match:
        return match.group(1)
    if re.fullmatch(r"[A-Z0-9]{10}", book):
        return book
    raise click.ClickException(f"Could not find an ASIN in {book!r}")


@cli.command()
@click.option(
    "--source",
    "source_path",
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Manuscript .docx from the ghostwriter",
)
@click.option(
    "--output",
    "output_path",
    required=True,
    type=click.Path(dir_okay=False, path_type=Path),
    help="Where to write the PDF report",
)
@click.option("--title", default=None, help="Report title (defaults to filename)")
@click.option(
    "--max-chapters",
    type=int,
    default=None,
    help="Limit to the first N chapters (quick preview / cost control)",
)
@click.option(
    "--claude-bin",
    default=DEFAULT_CLAUDE_BIN,
    show_default=True,
    help="Claude Code CLI binary",
)
@click.option(
    "--book",
    default=None,
    help="Amazon book link or ASIN (fetches reviews -> review mode)",
)
@click.option(
    "--reviews-file",
    "reviews_file",
    default=None,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="Pre-fetched reviews .txt (alternative to --book)",
)
@click.option(
    "--ghostwriter",
    default=None,
    help="Ghostwriter name (builds their track record; see 'scorecard')",
)
def analyze(
    source_path: Path,
    output_path: Path,
    title: str | None,
    max_chapters: int | None,
    claude_bin: str,
    book: str | None,
    reviews_file: Path | None,
    ghostwriter: str | None,
) -> None:
    """Analyse a manuscript's quality and write a PDF report.

    Pre-publication mode: a general craft / genre / consistency review with a
    page-1 verdict, top issues (each backed by a verbatim quote), a chapter
    heatmap and a ready-to-forward fix list.

    Review mode (--book / --reviews-file): also distils reader complaints
    and flags whether this manuscript repeats them (regression check).
    """
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    report_title = title or source_path.stem

    click.echo(f"Reading {source_path.name}...")
    doc = Document(str(source_path))
    paragraphs = [
        (para.style.name if para.style and para.style.name else "", para.text)
        for para in doc.paragraphs
    ]
    chapters = split_chapters(paragraphs)
    if max_chapters:
        chapters = chapters[:max_chapters]
    click.echo(f"Detected {len(chapters)} chapter(s).")

    all_text = [text for _, text in paragraphs]
    chapter_metrics = [
        ChapterMetrics(title, len(body.split())) for title, body in chapters
    ]
    metrics = compute_metrics(all_text, chapters=chapter_metrics)

    def runner(prompt: str) -> str:
        return claude_runner(prompt, claude_bin=claude_bin)

    reviews_text: str | None = None
    if reviews_file is not None:
        reviews_text = reviews_file.read_text(encoding="utf-8")
        click.echo(f"Loaded reviews from {reviews_file.name} (review mode)")
    elif book:
        asin = _extract_asin(book)
        click.echo(f"Fetching reviews for {asin} (review mode)...")
        reviews = fetch_reviews(asin)
        reviews_text = format_reviews_report(reviews, asin)
        click.echo(f"Fetched {len(reviews)} reviews")

    learning_text = load_learning()
    click.echo(
        f"Analysing with Claude ({len(chapters)} chapters + synthesis; "
        "this takes a few minutes)..."
    )
    analysis = analyze_manuscript(
        chapters, metrics, runner, reviews_text=reviews_text, learning=learning_text
    )

    write_pdf(analysis, output_path, title=report_title)
    syn = analysis.synthesis
    click.echo(f"Verdict: {syn.verdict.upper()} ({syn.score}/100)")
    click.echo(f"Report -> {output_path}")

    click.echo("Updating the learning overview...")
    save_learning(update_learning(learning_text, analysis, runner, book=report_title))

    save_record(record_from_analysis(analysis, report_title, ghostwriter or ""))
    if ghostwriter:
        click.echo(f"Recorded to {ghostwriter}'s track record (see 'scorecard').")


@cli.command()
def scorecard() -> None:
    """Show each ghostwriter's track record across past analyses."""
    cards = build_scorecard(load_history())
    if not cards:
        click.echo("No history yet. Run 'analyze --ghostwriter <name>' first.")
        return
    for card in cards:
        click.echo(
            f"\n{card.ghostwriter} - {card.book_count} book(s), "
            f"avg {card.avg_score}/100"
        )
        if card.weak_dimensions:
            weak = ", ".join(f"{dim} ({n})" for dim, n in card.weak_dimensions)
            click.echo(f"  weakest areas: {weak}")
        for book_title, score in card.books:
            click.echo(f"  - {book_title}: {score}/100")


if __name__ == "__main__":
    cli()
